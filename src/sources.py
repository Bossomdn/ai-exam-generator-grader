"""แปลงเนื้อหาที่ผู้ใช้ให้มา (PPTX / PDF / TXT / หัวข้อ) เป็น content block สำหรับ Claude"""

import base64
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Literal, Optional

import pypdfium2 as pdfium
from pptx import Presentation

SourceKind = Literal["pptx", "pdf", "text", "topic"]


@dataclass
class Source:
    kind: SourceKind
    name: str
    text: Optional[str] = None     # pptx / text / topic
    pdf_b64: Optional[str] = None  # pdf

    @property
    def is_topic_only(self) -> bool:
        return self.kind == "topic"

    def content_blocks(self) -> list[dict]:
        """เนื้อหาวางไว้ก่อนคำสั่ง และใส่ cache_control เพื่อให้เรียกซ้ำ (เช่นสร้างแบบฝึกเพิ่ม) ถูกลง"""
        cache = {"type": "ephemeral"}
        if self.kind == "pdf":
            return [{
                "type": "document",
                "source": {"type": "base64", "media_type": "application/pdf", "data": self.pdf_b64},
                "title": self.name,
                "cache_control": cache,
            }]
        if self.kind == "topic":
            return [{"type": "text", "text": f"<topic>\n{self.text}\n</topic>"}]
        return [{
            "type": "text",
            "text": f'<course_material name="{self.name}">\n{self.text}\n</course_material>',
            "cache_control": cache,
        }]


def _shape_texts(shape) -> list[str]:
    texts: list[str] = []
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        for child in shape.shapes:
            texts.extend(_shape_texts(child))
        return texts
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            line = "".join(run.text for run in paragraph.runs).strip()
            if line:
                texts.append(line)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                texts.append(" | ".join(cells))
    return texts


def load_pptx(data: bytes, name: str) -> Source:
    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for number, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_texts(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"(โน้ตผู้สอน) {notes}")
        body = "\n".join(lines) if lines else "(ไม่มีข้อความ)"
        parts.append(f"[สไลด์ {number}]\n{body}")
    return Source(kind="pptx", name=name, text="\n\n".join(parts))


def pdf_page_texts(data: bytes) -> list[str]:
    """ข้อความรายหน้า (pdfium ดึงสระ/วรรณยุกต์ไทยได้ครบกว่า pypdf)"""
    doc = pdfium.PdfDocument(data)
    try:
        return [doc[i].get_textpage().get_text_range().strip() for i in range(len(doc))]
    finally:
        doc.close()


def pdf_to_text(data: bytes) -> str:
    """ใช้กับโมเดลที่อ่าน PDF เองไม่ได้ (เช่น Ollama) — ได้เฉพาะข้อความ ไม่เห็นรูป/ตาราง"""
    return "\n\n".join(
        f"[หน้า {number}]\n{text or '(ไม่มีข้อความ)'}"
        for number, text in enumerate(pdf_page_texts(data), start=1)
    )


def load_pdf(data: bytes, name: str) -> Source:
    return Source(kind="pdf", name=name, pdf_b64=base64.standard_b64encode(data).decode("ascii"))


def load_text(data: bytes, name: str) -> Source:
    return Source(kind="text", name=name, text=data.decode("utf-8-sig"))


def from_topic(topic: str) -> Source:
    return Source(kind="topic", name=topic.strip(), text=topic.strip())


def load_upload(name: str, data: bytes) -> Source:
    suffix = Path(name).suffix.lower()
    if suffix == ".pptx":
        return load_pptx(data, name)
    if suffix == ".pdf":
        return load_pdf(data, name)
    if suffix in (".txt", ".md"):
        return load_text(data, name)
    raise ValueError(f"ยังไม่รองรับไฟล์ {suffix} (รองรับ .pptx .pdf .txt .md)")
