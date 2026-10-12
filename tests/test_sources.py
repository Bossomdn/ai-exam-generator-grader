import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from src import sources


def make_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = "Overfitting"
    slide.placeholders[1].text = "โมเดลจำข้อมูลฝึกมากเกินไป"
    slide.notes_slide.notes_text_frame.text = "ยกตัวอย่างการสอบท่องจำ"

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide2.shapes.title.text = "Metrics"
    table = slide2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Precision"
    table.cell(0, 1).text = "TP / (TP + FP)"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pptx_extracts_text_per_slide_with_notes_and_tables():
    src = sources.load_upload("lecture.pptx", make_pptx())
    assert src.kind == "pptx"
    assert "[สไลด์ 1]" in src.text and "[สไลด์ 2]" in src.text
    assert "โมเดลจำข้อมูลฝึกมากเกินไป" in src.text
    assert "(โน้ตผู้สอน) ยกตัวอย่างการสอบท่องจำ" in src.text
    assert "Precision | TP / (TP + FP)" in src.text


def test_text_source_strips_bom_and_is_cached():
    src = sources.load_upload("notes.txt", "﻿เนื้อหา".encode("utf-8"))
    block = src.content_blocks()[0]
    assert block["type"] == "text"
    assert "<course_material" in block["text"] and "﻿" not in block["text"]
    assert block["cache_control"] == {"type": "ephemeral"}


def test_pdf_source_becomes_document_block():
    src = sources.load_upload("slides.pdf", b"%PDF-1.4 fake")
    block = src.content_blocks()[0]
    assert block["type"] == "document"
    assert block["source"]["media_type"] == "application/pdf"


def test_topic_only_source():
    src = sources.from_topic("  ระบบสุริยะ ")
    assert src.is_topic_only
    assert src.content_blocks()[0]["text"] == "<topic>\nระบบสุริยะ\n</topic>"


def test_unsupported_file_type():
    with pytest.raises(ValueError):
        sources.load_upload("data.xlsx", b"")
